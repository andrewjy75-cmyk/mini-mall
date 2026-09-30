"""文档解析：把 PDF / Word / TXT 变成纯文本。

这个文件看起来没什么技术含量，但它是整条链路的**质量天花板**：
解析阶段丢掉的信息，后面的切分、向量化、重排再怎么优化都找不回来——
垃圾进，垃圾出。

这里踩过的坑，每一个都会让最终效果明显变差，所以都写在代码里了：

坑 1｜扫描版 PDF 提不出文字
    很多"红头文件"、印刷体政策文档是扫描件，PDF 里根本没有文字层，
    pypdf 提取出来是空字符串。如果不管，就会往知识库里塞一堆空文档，
    检索时永远命中不了，排查起来极其痛苦（因为文档列表里"看起来"是成功的）。
    → 所以提取完必须做"空内容检查"，宁可报错也不静默入库。

坑 2｜PDF 提取出来的换行是"视觉换行"
    PDF 里每一行文字的位置是独立记录的，提取出来就是一行一个碎片：
        "退换货政策如下：\n商品签收后 7 天内\n可无理由退货。"
    这些 \n 是排版需要，不是句子的结束。直接连起来会把一句话拆散，
    影响后续切分和向量化的语义完整性。
    → 所以要智能地把"被硬换行打断的句子"接回去。

坑 3｜Word 的表格不在段落里
    python-docx 的 doc.paragraphs 【不包含】表格里的文字。
    而商品规格、运费标准这类关键信息恰恰最爱用表格。
    只读 paragraphs 会静默丢掉最重要的那部分内容。
    → 必须单独遍历 doc.tables，并转成 Markdown 表格保留行列结构。
"""

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from langchain_core.documents import Document

logger = logging.getLogger(__name__)

# 支持的扩展名。加新格式时只改这里 + _PARSERS 注册表
SUPPORTED_EXTENSIONS: set[str] = {".pdf", ".docx", ".txt", ".md"}

# 自定义异常：解析失败要能明确区分于"文件不存在"这类系统错误，
# 上层好决定是提示用户重新上传，还是记日志报警
class DocumentParseError(Exception):
    """文档解析失败。message 里是给用户看的人话。"""


@dataclass
class ParseOptions:
    """解析选项。集中放在一个对象里，以后加选项不用改函数签名。"""

    # 是否保留页码信息（PDF 有，其他格式没有）
    keep_page_number: bool = True
    # 解析出的正文少于这个字符数，就判定为"没提到有效内容"
    min_chars: int = 20
    # 额外元数据（如文档 id、标题），会原样挂到每个 Document 上
    extra_metadata: dict = field(default_factory=dict)


# ==================== 文本清洗 ====================

# 零宽字符：从网页复制到 PDF 时经常夹带，肉眼看不见但会污染向量
_ZERO_WIDTH = re.compile(r"[\u200b-\u200f\u2028\u2029\ufeff]")
# 连续空白（含全角空格）折叠成一个空格
_MULTI_SPACE = re.compile(r"[ \t\u3000]{2,}")
# 3 个以上连续换行压缩成 2 个（保留段落感，去掉大片空白）
_MULTI_NEWLINE = re.compile(r"\n{3,}")
# 中文字符区间。用于识别"中英混排"的边界
_CJK = r"\u4e00-\u9fff\u3000-\u303f\uff00-\uffef"

# ---- 判断「PDF 换行该怎么处理」的三个信号 ----
# 详见 _join_wrapped_lines 的说明。这三个正则一起决定了
# 哪些换行是排版产生（该合并）、哪些是结构边界（不该合并）

# 信号 1｜句末标点：上一行以这些结尾，说明话说完了，不该再往下接
_SENTENCE_END = re.compile(r"[。！？；：!?;:]\s*$")

# 信号 2｜被截断的尾巴：上一行以数字或字母结尾，说明词/数字被切开了。
# 例：'自商品签收之日起 7' + '个自然日内…' —— 这行很短，但必须合并
_CONTINUATION_TAIL = re.compile(r"[0-9A-Za-z]\s*$")

# 信号 3｜新块开头：下一行以列表符号、序号或标题记号开头，说明它是新的一项。
# 没有这条，连续的列表项会被粘成一长串（这是很常见的误合并）
_NEW_BLOCK_HEAD = re.compile(
    r"^\s*(?:[·•▪◦\-*]|#+|\d+\s*[、.)）]"
    r"|[一二三四五六七八九十]+\s*[、.．)）]"
    r"|第[一二三四五六七八九十百零〇]+\s*[条章节款项])"
)


def _normalize_text(text: str) -> str:
    """通用清洗：去零宽字符、统一换行、压缩多余空白。"""
    if not text:
        return ""
    text = _ZERO_WIDTH.sub("", text)
    # PDF 提取的换行可能是 \r\n / \r，统一成 \n
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _MULTI_SPACE.sub(" ", text)
    text = _MULTI_NEWLINE.sub("\n\n", text)
    return text.strip()


def _join_wrapped_lines(text: str) -> str:
    """把 PDF 的「视觉换行」接回完整句子（坑 2）。

    ★ 这一版规则是踩过坑之后重写的，先说清楚上一版错在哪：

      上一版的规则是「上一行末尾是中文、下一行开头是中文 → 合并」。
      听起来合理，但实测把标题和正文粘成了一段：

          'Mini Mall 退换货政策'            ← 标题，15 字
          '生效日期：2026 年 1 月 1 日。…'   ← 正文，44 字
          合并结果：'Mini Mall 退换货政策生效日期：…'   ← 标题没了

      问题在于：PDF 里有两类完全不同的换行，光看「是不是中文」区分不出来。
          A. 排版换行：这一行被排满了，句子还没说完   → 要合并
          B. 结构换行：标题结束、段落结束、列表项结束 → 不能合并

      怎么区分？抓住两个可观测的特征：

      1) 排版换行会让上一行【接近满行】。
         上面例子里，正文行 44 字、标题行 15 字，差距很明显。
         所以先用整页行长度的 75 分位数估算「行宽」，
         只有长度接近行宽的上一行，才可能是被排满截断的。
      2) 还有一种短的续行：上一行以数字或字母结尾，
         说明这个词/数字被切断了（例：'自商品签收之日起 7' + '个自然日内…'）。
         这种虽然短，也该合并。

      再叠加两条否决条件，覆盖剩下的边界：
      3) 上一行以句末标点结尾（。！？；：）→ 话说完了，不合并。
      4) 下一行以列表符号或序号开头（·、-、一、、第二条…）→ 它是新的一项，不合并。
         这条很关键：连续的列表项之间如果不拦住，会被粘成一长串。

    这套规则在本项目的 PDF 样本上逐行验证通过（标题、正文、续行、列表四种情况都对）。
    但它终究是启发式规则，不是 PDF 解析的通用解——
    更稳的方案是换用带版面分析能力的库（如 pdfplumber），
    或者直接用结构化格式（docx）。
    """
    lines = text.split("\n")
    if len(lines) <= 1:
        return text

    # ---- 估算本页的「行宽」----
    # 用 75 分位数而不是最大值：最大值可能是个别超长行（比如没有空格的表格行），
    # 用它会导致阈值虚高，正常的满行反而被判成「没排满」
    lengths = sorted(len(line.strip()) for line in lines if line.strip())
    if lengths:
        width_hint = lengths[int(len(lengths) * 0.75)]
    else:
        width_hint = 0
    # 下限 20 字：极短的文档（比如只有几行）不该因为阈值太小而疯狂合并
    near_full_width = max(width_hint * 0.8, 20)

    merged: list[str] = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            # 空行是段落分隔，必须保留——它是切分器判断段落边界的重要依据
            merged.append("")
            continue

        if merged and merged[-1]:
            prev = merged[-1]
            should_merge = (
                # 条件 3：上一行没说完整
                not _SENTENCE_END.search(prev)
                # 条件 4：下一行不是新的一项
                and not _NEW_BLOCK_HEAD.match(stripped)
                # 条件 1 或 2：上一行接近满行，或者结尾是被截断的数字/字母
                and (
                    len(prev) >= near_full_width
                    or bool(_CONTINUATION_TAIL.search(prev))
                )
            )

            # 英文断词单独处理：'re-\nview' → 'review'
            # 这个优先级最高，因为连字符结尾是很强的断词信号
            if prev.endswith("-") and re.match(r"^[a-z]", stripped):
                merged[-1] = prev[:-1] + stripped
                continue

            if should_merge:
                merged[-1] = prev + stripped
                continue

        merged.append(stripped)

    return "\n".join(merged)


# ==================== 各格式解析器 ====================


def _parse_pdf(path: Path, options: ParseOptions) -> list[Document]:
    """解析 PDF：逐页提取，每页一个 Document。

    逐页而不是整篇一个 Document，是为了保留页码。
    用户看到回答时会问"这是哪一页写的"，能答上来才叫溯源。
    """
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
    except (PdfReadError, OSError, ValueError) as exc:
        raise DocumentParseError(f"PDF 文件无法打开，可能已损坏：{exc}") from exc

    # PDF 加密（有些文档只设了"禁止复制"密码）
    if reader.is_encrypted:
        try:
            # 先试空密码——很多 PDF 只是设了权限位，并没有真正的打开密码
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001
            raise DocumentParseError("该 PDF 已加密，请先解除密码后再上传") from exc

    docs: list[Document] = []
    empty_pages = 0

    for page_no, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001
            # 单页失败不该让整份文档失败：继续解析其余页，只记日志
            logger.warning("PDF 第 %s 页解析失败，已跳过：%s", page_no, exc)
            continue

        text = _join_wrapped_lines(_normalize_text(raw))
        if len(text) < 10:
            empty_pages += 1
            continue

        metadata = {"source": path.name, "page": page_no, "file_type": "pdf"}
        if options.keep_page_number:
            metadata["page_label"] = f"第 {page_no} 页"
        metadata.update(options.extra_metadata)
        docs.append(Document(page_content=text, metadata=metadata))

    # ★ 坑 1 的处理：一页都没提出文字，基本可以断定是扫描件
    if not docs:
        raise DocumentParseError(
            "未能从该 PDF 中提取到任何文字。"
            "如果它是扫描件或纯图片 PDF，需要先做 OCR 识别——"
            "本项目暂未内置 OCR，请改用可复制文字的版本。"
        )
    if empty_pages:
        logger.info("PDF %s 有 %s 页无文字（可能是插图页），已跳过", path.name, empty_pages)

    return docs


def _parse_docx(path: Path, options: ParseOptions) -> list[Document]:
    """解析 Word：段落按顺序抽出，表格转成 Markdown（坑 3）。

    关于"保持段落顺序"：docx 的文档体是一个元素序列，
    段落和表格是交错排列的。如果先把所有段落读完再读所有表格，
    表格就会跑到文末去，上下文顺序全乱。
    所以这里按 body 里的真实顺序遍历。
    """
    import docx
    from docx.document import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document: DocxDocument = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"Word 文件无法打开，可能已损坏：{exc}") from exc

    blocks: list[str] = []

    def _iter_body_items(doc):
        """按文档真实顺序，依次产出段落和表格。"""
        body = doc.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, doc)
            elif child.tag == qn("w:tbl"):
                yield Table(child, doc)

    for item in _iter_body_items(document):
        if isinstance(item, Paragraph):
            text = _normalize_text(item.text)
            if not text:
                continue
            # 保留标题层级：切分时"标题"是天然的语义边界，
            # 而且检索时标题里的词往往是命中关键
            style_name = (item.style.name or "") if item.style else ""
            if style_name.startswith("Heading") or style_name.startswith("标题"):
                level = "".join(ch for ch in style_name if ch.isdigit()) or "1"
                blocks.append(f"{'#' * min(int(level), 6)} {text}")
            else:
                blocks.append(text)

        elif isinstance(item, Table):
            markdown = _table_to_markdown(item)
            if markdown:
                blocks.append(markdown)

    text = _join_wrapped_lines(_normalize_text("\n\n".join(blocks)))
    if len(text) < options.min_chars:
        raise DocumentParseError(
            "该 Word 文档里没有提取到有效文字内容，请检查文件是否为空。"
        )

    metadata = {"source": path.name, "page": None, "file_type": "docx"}
    metadata.update(options.extra_metadata)
    return [Document(page_content=text, metadata=metadata)]


def _table_to_markdown(table) -> str:
    """把 Word 表格转成 Markdown 表格。

    为什么要转成 Markdown 而不是直接拼文字？
    因为拼成 "商品名 价格 库存" 这种一行串，模型根本分不清哪个是列名、
    哪一行对应哪一列，回答"某商品多少钱"时就容易串行（张冠李戴的幻觉）。
    Markdown 的表头 + 分隔线能把行列结构明确告诉模型。
    """
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [_normalize_text(cell.text).replace("\n", " ") for cell in row.cells]
        if any(cells):
            rows.append(cells)

    if not rows:
        return ""

    width = max(len(r) for r in rows)
    # 补齐每行列数，否则 Markdown 表格渲染时会错位
    rows = [r + [""] * (width - len(r)) for r in rows]

    lines = ["| " + " | ".join(rows[0]) + " |"]
    lines.append("| " + " | ".join(["---"] * width) + " |")
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _parse_text(path: Path, options: ParseOptions) -> list[Document]:
    """解析纯文本 / Markdown。

    编码问题是这里的唯一坑：中文 Windows 上导出的 txt 可能是 GBK，
    按 UTF-8 读会直接抛 UnicodeDecodeError。
    """
    raw_bytes = path.read_bytes()
    text = ""
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "gbk"):
        try:
            text = raw_bytes.decode(encoding)
            logger.debug("以 %s 成功解码 %s", encoding, path.name)
            break
        except UnicodeDecodeError:
            continue
    else:
        # errors="ignore" 是最后的兜底：宁可丢几个乱码字符，也别让上传失败
        text = raw_bytes.decode("utf-8", errors="ignore")
        logger.warning("%s 编码无法识别，已按 UTF-8 容错解码", path.name)

    text = _join_wrapped_lines(_normalize_text(text))
    if len(text) < options.min_chars:
        raise DocumentParseError("文本文件内容为空或过短，没有提取价值。")

    metadata = {
        "source": path.name,
        "page": None,
        "file_type": path.suffix.lstrip(".").lower(),
    }
    metadata.update(options.extra_metadata)
    return [Document(page_content=text, metadata=metadata)]


# 注册表：扩展名 → 解析函数。加格式只动这一行
_PARSERS = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".txt": _parse_text,
    ".md": _parse_text,
}


# ==================== 对外入口 ====================


def is_supported(filename: str) -> bool:
    """这个文件能不能解析。上传接口先调它做校验，比等解析失败再报错体验好。"""
    return Path(filename).suffix.lower() in SUPPORTED_EXTENSIONS


def supported_extensions_text() -> str:
    """给人看的支持格式说明，用于错误提示。"""
    return " / ".join(sorted(e.lstrip(".").upper() for e in SUPPORTED_EXTENSIONS))


def load_document(
    path: str | Path, options: ParseOptions | None = None
) -> list[Document]:
    """把一个文件解析成 Document 列表（PDF 每页一条，其他整篇一条）。

    参数:
        path: 文件路径
        options: 解析选项，不传就用默认

    异常:
        DocumentParseError: 格式不支持、文件损坏、或没提取到有效内容
    """
    path = Path(path)
    options = options or ParseOptions()

    if not path.exists():
        raise DocumentParseError(f"文件不存在：{path}")

    suffix = path.suffix.lower()
    parser = _PARSERS.get(suffix)
    if parser is None:
        raise DocumentParseError(
            f"暂不支持 {suffix or '该'} 格式，目前支持：{supported_extensions_text()}"
        )

    docs = parser(path, options)
    total_chars = sum(len(d.page_content) for d in docs)
    logger.info(
        "解析完成 %s：%s 段，共 %s 字符", path.name, len(docs), total_chars
    )
    return docs
